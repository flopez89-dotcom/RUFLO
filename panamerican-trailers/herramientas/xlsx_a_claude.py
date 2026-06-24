#!/usr/bin/env python3
"""
xlsx_a_claude.py — Convierte un Excel (.xlsx) a Markdown listo para pegar
en Claude (claude.ai o Claude Code), evitando el add-in roto.

Uso:
    python3 xlsx_a_claude.py archivo.xlsx
    python3 xlsx_a_claude.py archivo.xlsx -o salida.md
    python3 xlsx_a_claude.py archivo.xlsx --max-filas 200

Genera un .md con:
- Resumen (nombre del archivo, hojas, dimensiones de cada hoja).
- Una sección por hoja con la tabla en formato Markdown.
- Si la hoja excede --max-filas, muestra primeras N + últimas 5 + total.
"""

from pathlib import Path
import argparse
import sys

try:
    import openpyxl
except ImportError:
    print("Falta openpyxl. Instala con: pip install openpyxl", file=sys.stderr)
    sys.exit(1)


def cell_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).replace("|", "\\|").replace("\n", " ")


def sheet_to_markdown(ws, max_rows: int) -> str:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return "_(hoja vacía)_\n"

    total = len(rows)
    header = rows[0]
    body = rows[1:]

    if total - 1 > max_rows:
        shown = body[:max_rows]
        tail = body[-5:] if len(body) > max_rows + 5 else []
    else:
        shown = body
        tail = []

    lines = []
    lines.append("| " + " | ".join(cell_str(c) for c in header) + " |")
    lines.append("| " + " | ".join("---" for _ in header) + " |")
    for r in shown:
        lines.append("| " + " | ".join(cell_str(c) for c in r) + " |")
    if tail:
        lines.append("| " + " | ".join("…" for _ in header) + " |")
        for r in tail:
            lines.append("| " + " | ".join(cell_str(c) for c in r) + " |")
        lines.append(f"\n_Mostrando primeras {max_rows} y últimas {len(tail)} de {total - 1} filas de datos._")
    else:
        lines.append(f"\n_Total: {total - 1} filas de datos._")
    return "\n".join(lines) + "\n"


def main():
    ap = argparse.ArgumentParser(description="Convierte un Excel a Markdown para Claude.")
    ap.add_argument("archivo", help="Ruta al archivo .xlsx")
    ap.add_argument("-o", "--output", help="Archivo .md de salida (por defecto: mismo nombre + .md)")
    ap.add_argument("--max-filas", type=int, default=500, help="Máximo de filas por hoja (default 500)")
    args = ap.parse_args()

    src = Path(args.archivo)
    if not src.exists():
        print(f"No existe: {src}", file=sys.stderr)
        sys.exit(1)

    out = Path(args.output) if args.output else src.with_suffix(".md")

    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)

    parts = []
    parts.append(f"# {src.name}\n")
    parts.append(f"_Hojas: {len(wb.sheetnames)} → {', '.join(wb.sheetnames)}_\n")

    for name in wb.sheetnames:
        ws = wb[name]
        parts.append(f"\n## Hoja: {name}\n")
        parts.append(f"_Dimensiones: {ws.max_row} filas x {ws.max_column} columnas_\n\n")
        parts.append(sheet_to_markdown(ws, args.max_filas))

    out.write_text("\n".join(parts), encoding="utf-8")

    size_kb = out.stat().st_size / 1024
    print(f"✓ Generado: {out}  ({size_kb:.1f} KB)")
    print(f"  Hojas: {len(wb.sheetnames)}")
    print("\nSiguiente paso:")
    print("  1. Abre el .md generado.")
    print("  2. Copia su contenido completo (Cmd/Ctrl+A, Cmd/Ctrl+C).")
    print("  3. Pégalo en claude.ai junto con tu pregunta.")


if __name__ == "__main__":
    main()
